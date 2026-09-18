"""Reads/writes the expense report workbook, matching the layout of the
user's own "Expense Reimbursement Form" template: dated rows 13-30 with
Date / Description / From / To / Mileage / Amount columns, and a
SUM formula total in row 31.
"""
import shutil
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "data" / "expense_report_template.xlsx"
INSTANCE_DIR = BASE_DIR / "instance"
REPORT_PATH = INSTANCE_DIR / "expense_report.xlsx"

SHEET_NAME = "Sheet1"
FIRST_DATA_ROW = 13
LAST_DATA_ROW = 30
COL_DATE = "A"
COL_DESCRIPTION = "B"
COL_FROM = "J"
COL_TO = "K"
COL_MILEAGE = "L"
COL_AMOUNT = "O"


class ReportFullError(Exception):
    """Raised when every row in the expense table already has data."""


def ensure_report_exists() -> Path:
    INSTANCE_DIR.mkdir(exist_ok=True)
    if not REPORT_PATH.exists():
        shutil.copyfile(TEMPLATE_PATH, REPORT_PATH)
    return REPORT_PATH


def reset_report():
    """Start a fresh report from the blank template."""
    INSTANCE_DIR.mkdir(exist_ok=True)
    shutil.copyfile(TEMPLATE_PATH, REPORT_PATH)


def _find_next_empty_row(ws) -> int:
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        if ws[f"{COL_DATE}{row}"].value in (None, ""):
            return row
    raise ReportFullError(
        f"The expense report is full ({LAST_DATA_ROW - FIRST_DATA_ROW + 1} rows). "
        "Start a new report before adding more expenses."
    )


def add_expense(date, description: str, amount: float,
                 from_addr: str = None, to_addr: str = None,
                 mileage: float = None) -> int:
    """Appends one expense row to the report. Returns the row number used."""
    ensure_report_exists()
    wb = openpyxl.load_workbook(REPORT_PATH)
    ws = wb[SHEET_NAME]
    row = _find_next_empty_row(ws)

    template_row = FIRST_DATA_ROW
    date_cell = ws[f"{COL_DATE}{row}"]
    date_cell.value = date
    date_cell.number_format = ws[f"{COL_DATE}{template_row}"].number_format

    ws[f"{COL_DESCRIPTION}{row}"] = description

    if from_addr:
        ws[f"{COL_FROM}{row}"] = from_addr
    if to_addr:
        ws[f"{COL_TO}{row}"] = to_addr
    if mileage is not None:
        ws[f"{COL_MILEAGE}{row}"] = mileage

    amount_cell = ws[f"{COL_AMOUNT}{row}"]
    amount_cell.value = amount
    amount_cell.number_format = ws[f"{COL_AMOUNT}{template_row}"].number_format

    wb.save(REPORT_PATH)
    return row


def list_expenses() -> dict:
    ensure_report_exists()
    wb = openpyxl.load_workbook(REPORT_PATH)
    ws = wb[SHEET_NAME]

    name = ws["B5"].value
    rows = []
    total = 0.0
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        date_val = ws[f"{COL_DATE}{row}"].value
        if date_val in (None, ""):
            continue
        amount_val = ws[f"{COL_AMOUNT}{row}"].value or 0
        total += float(amount_val)
        rows.append({
            "row": row,
            "date": date_val.strftime("%m/%d/%Y") if hasattr(date_val, "strftime") else date_val,
            "description": ws[f"{COL_DESCRIPTION}{row}"].value,
            "from_addr": ws[f"{COL_FROM}{row}"].value,
            "to_addr": ws[f"{COL_TO}{row}"].value,
            "mileage": ws[f"{COL_MILEAGE}{row}"].value,
            "amount": amount_val,
        })
    return {"name": name, "entries": rows, "total": round(total, 2)}
